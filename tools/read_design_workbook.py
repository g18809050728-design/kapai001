# -*- coding: utf-8 -*-
"""
读回《设计元素表》Excel，只报告「你的决定 与 当前值 不同」的项。

这是助手了解人类决策的入口：不需要通读全表，只看本脚本的输出即可。
运行： python tools/read_design_workbook.py
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "docs", "设计元素表.xlsx")

# 每张值表：id 列 / 名称列 / 当前值列 / 决定列
SHEETS = [
    # 玩法侧
    ("游戏基本信息", 1, 2, [3], [4], ["值"], ["值"]),
    ("实体信息", 1, 2, [4], [5], ["值"], ["值"]),
    ("卡牌信息", 1, 2, [4, 7, 6], [9, 10, 11], ["费用", "数量", "效果"], ["费用", "数量", "效果"]),
    ("游戏机制", 1, 2, [3], [4], ["规则"], ["规则"]),
    # 表现侧
    ("界面", 1, 2, [3], [4], ["值"], ["值"]),
    ("色彩", 1, 2, [3], [4], ["色值"], ["色值"]),
    ("卡牌类型色", 1, 2, [3], [4], ["色值"], ["色值"]),
    ("字号", 1, 2, [3], [4], ["字号"], ["字号"]),
    ("元素尺寸", 1, 2, [3, 4], [5, 6], ["宽", "高"], ["宽", "高"]),
    ("布局", 1, 2, [3, 4, 5, 6], [7, 8, 9, 10], ["x", "y", "w", "h"], ["x", "y", "w", "h"]),
    ("美术资源", 1, 2, [2], [5], ["路径"], ["路径"]),
]

# 这些填写内容视为「不改」
NO_CHANGE = {"", "保持", "不变", "-", "--", "keep", "same"}


def norm(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s.lower() in NO_CHANGE:
            return None
        return s
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def fmt(v):
    return "（空）" if v is None else str(v)


def main():
    if not os.path.exists(XLSX):
        print("找不到 %s" % XLSX)
        return 1
    import datetime
    mt = datetime.datetime.fromtimestamp(os.path.getmtime(XLSX)).strftime("%Y-%m-%d %H:%M")
    wb = load_workbook(XLSX, data_only=True)

    print("=" * 70)
    print("设计元素表 · 决策读取")
    print("文件: docs/设计元素表.xlsx    最后修改: %s" % mt)
    print("=" * 70)

    total = 0
    for name, idc, labc, curc, decc, cur_n, dec_n in SHEETS:
        if name not in wb.sheetnames:
            print("\n[%s]  ! 工作表不存在，跳过" % name)
            continue
        ws = wb[name]
        rows_all = 0
        changes = []
        for r in range(3, ws.max_row + 1):
            rid = norm(ws.cell(row=r, column=idc).value)
            if rid is None:
                continue
            rows_all += 1
            label = fmt(norm(ws.cell(row=r, column=labc).value))
            cur = [norm(ws.cell(row=r, column=c).value) for c in curc]
            dec = [norm(ws.cell(row=r, column=c).value) for c in decc]
            if all(d is None for d in dec):
                continue
            if len(cur) == 1 and len(dec) == 1:
                if dec[0] != cur[0]:
                    changes.append((rid, label, "%s -> %s" % (fmt(cur[0]), fmt(dec[0]))))
            else:
                diffs = []
                for i in range(min(len(cur), len(dec))):
                    if dec[i] is not None and dec[i] != cur[i]:
                        diffs.append("%s: %s -> %s" % (dec_n[i], fmt(cur[i]), fmt(dec[i])))
                if diffs:
                    changes.append((rid, label, "；".join(diffs)))

        if changes:
            total += len(changes)
            print("\n[%s]  已决定 %d 项 / 共 %d 项" % (name, len(changes), rows_all))
            w = max(len(c[0]) for c in changes) + 2
            for rid, label, desc in changes:
                print("   %-*s %-20s %s" % (w, rid, label[:20], desc))
        else:
            print("\n[%s]  未改动（共 %d 项）" % (name, rows_all))

    print("\n" + "=" * 70)
    if total == 0:
        print("汇总：没有任何决定 —— 全部保持当前实现。")
    else:
        print("汇总：共 %d 项决定，请按上面的清单修改代码。" % total)
    print("=" * 70)
    return 0


if __name__ == "__main__":
    sys.exit(main())