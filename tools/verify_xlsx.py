# -*- coding: utf-8 -*-
"""读回生成的 xlsx，验证结构是否完整。"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from openpyxl import load_workbook

PATH = os.path.join("docs", "PVE卡牌Demo_升级方案.xlsx")
wb = load_workbook(PATH)
print("文件:", os.path.abspath(PATH))
print("工作表数:", len(wb.sheetnames))
print()

problems = []
total_rows = 0
for name in wb.sheetnames:
    ws = wb[name]
    title = ws.cell(row=1, column=1).value
    headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    headers = [h for h in headers if h is not None]
    data_rows = 0
    empty_rows = []
    for r in range(3, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        if all(v is None or str(v).strip() == "" for v in vals):
            empty_rows.append(r)
        else:
            data_rows += 1
    total_rows += data_rows
    print("── %s ──" % name)
    print("   标题   : %s" % title)
    print("   表头   : %s" % " | ".join(str(h) for h in headers))
    print("   数据行 : %d   冻结:%s   筛选:%s" % (data_rows, ws.freeze_panes, ws.auto_filter.ref))
    # 首行数据抽样
    print("   首行   : %s" % " | ".join(str(ws.cell(row=3, column=c).value)[:28] for c in range(1, len(headers) + 1)))
    print()

    if title is None:
        problems.append("%s: 缺标题" % name)
    if len(headers) < 2:
        problems.append("%s: 表头异常" % name)
    if data_rows == 0:
        problems.append("%s: 无数据" % name)
    if ws.freeze_panes is None:
        problems.append("%s: 未冻结表头" % name)

print("=" * 60)
print("总数据行:", total_rows)
if problems:
    print("发现问题:")
    for p in problems:
        print("  -", p)
else:
    print("结构检查: 全部通过")
