"""Check design delivery files; never import or execute the game."""
from collections import Counter
from datetime import datetime
import hashlib
import json
from pathlib import Path
import platform
import re
import runpy
import sys
from urllib.parse import unquote
import zipfile

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
COMMAND = 'python artifacts/redesign-20261030/verify_delivery.py'
FIELDS = ['id', 'owner', 'character', 'name', 'name_key', 'group', 'cost',
          'upgraded_cost', 'type', 'rarity', 'base', 'upgraded', 'tags',
          'permanent', 'exam_condition', 'exam', 'exam_upgraded', 'chapter',
          'chapter_upgraded']
RESULTS = []


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read(path):
    return (ROOT / path).read_text(encoding='utf-8')


def check(name, action):
    try:
        detail = action()
        result = {'check': name, 'passed': True, 'detail': detail}
    except Exception as exc:
        result = {'check': name, 'passed': False,
                  'detail': f'{type(exc).__name__}: {exc}'}
    RESULTS.append(result)
    print(('PASS ' if result['passed'] else 'FAIL ') + name)
    if not result['passed']:
        print(result['detail'])


def require(condition, detail):
    if not condition:
        raise AssertionError(detail)


def flat(card):
    return ['、'.join(card[k]) if isinstance(card[k], list) else card[k]
            for k in FIELDS]


def normalized(rows):
    return [[v if v is not None else '' for v in row] for row in rows]


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    # Reports are linked by the delivery index and replaced below with real results.
    # A first run must be able to resolve those output paths before writing its report.
    pending_reports = {
        'validation.md': '# 核查进行中\n\n本文件将在本次检查结束后写入实际结果，当前不代表通过。\n',
        'validation.json': '{"status": "verification in progress; no pass claimed"}\n',
    }
    for name, pending in pending_reports.items():
        if not (HERE / name).exists():
            (HERE / name).write_text(pending, encoding='utf-8')
    data = json.loads((HERE / 'campus-cards-v3.json').read_text(encoding='utf-8'))
    baseline = json.loads((HERE / 'baseline.json').read_text(encoding='utf-8'))
    protected = json.loads((HERE / 'protected-before.json').read_text(encoding='utf-8'))
    revision_dir = HERE / 'debuff-revision'
    revision_before = json.loads((revision_dir / 'before-v3.1.json').read_text(encoding='utf-8'))
    cards, tokens = data['cards'], data['tokens']
    all_cards = cards + tokens
    by_id = {c['id']: c for c in all_cards}
    guides = runpy.run_path(str(HERE / 'build_delivery.py'))['GUIDES']
    markdown = read('docs/校园卡牌设计表_v3.md')
    book = openpyxl.load_workbook(HERE / '校园角色与专属卡池_v3.xlsx', data_only=False)

    def counts():
        require(len(cards) == 172 and len(tokens) == 2, '172 permanent + 2 tokens required')
        owners = Counter(c['owner'] for c in cards)
        require(owners == {'C': 12, 'P01': 40, 'P02': 40, 'P03': 40, 'P04': 40}, owners)
        for owner in ('P01', 'P02', 'P03', 'P04'):
            own = [c for c in cards if c['owner'] == owner]
            groups = Counter(c['group'] for c in own)
            rarity = Counter(c['rarity'] for c in own)
            require(len(groups) == 4 and set(groups.values()) == {10}, (owner, groups))
            require(rarity == {'普通': 20, '进阶': 15, '稀有': 5}, (owner, rarity))
        return dict(owners)
    check('card_counts_groups_and_rarities', counts)

    def definitions():
        for key in ('id', 'name', 'name_key'):
            require(len({c[key] for c in all_cards}) == 174, f'Duplicate {key}')
        for c in all_cards:
            require(set(c) == set(FIELDS), (c['id'], 'field mismatch'))
            require(c['character'] == data['owners'][c['owner']], c['id'])
            require(c['type'] in {'工作', '策略', '能力', '合作', '场景'}, c['id'])
            require(type(c['cost']) is int and type(c['upgraded_cost']) is int
                    and 0 <= c['upgraded_cost'] <= c['cost'], (c['id'], 'cost'))
            for k in ('name', 'name_key', 'group', 'base', 'upgraded'):
                require(isinstance(c[k], str) and c[k].strip(), (c['id'], k))
                require('待补' not in c[k] and 'TODO' not in c[k], (c['id'], k))
            require(bool(c['tags']), (c['id'], 'tags'))
            require(c['permanent'] == (c in cards), (c['id'], 'lifecycle'))
        source = runpy.run_path(str(HERE / 'catalog_v3.py'))
        require(source['CARDS'] == cards and source['TOKENS'] == tokens,
                'Hand-authored catalog differs from JSON')
        require(source['STARTERS'] == data['starters'], 'Starter source mismatch')
        require(source['DEBUFFS'] == data['debuffs'], 'Debuff source mismatch')
        require(source['FEATURE_COVERAGE'] == data['feature_coverage'], 'Feature source mismatch')
        refs = set()
        for c in all_cards:
            for key in ('base', 'upgraded', 'exam', 'exam_upgraded', 'chapter', 'chapter_upgraded'):
                refs.update(re.findall(r'R3-(?:[CP]\d{2}-\d{2}|C-\d{2}|T\d{3}|[US]\d{2})', c[key]))
        allowed = set(by_id) | {'R3-U01', 'R3-U02', 'R3-U03', 'R3-S01', 'R3-S02'}
        require(not refs - allowed, sorted(refs - allowed))
        return {'unique_definitions': 174, 'source_json_match': True,
                'card_text_references_checked': len(refs)}
    check('definition_completeness_and_references', definitions)

    def faces():
        exams = [c for c in cards if c['exam']]
        chapters = [c for c in cards if c['chapter']]
        require(len(exams) == 14 and len(chapters) == 6, '14 exam / 6 chapter faces')
        require(Counter(c['type'] for c in exams) == {'合作': 7, '策略': 5, '场景': 2},
                'Exam face type coverage')
        require(Counter(c['exam_condition'] for c in exams) == {'独立作答': 12, '固定考场': 2},
                'Exam condition coverage')
        for c in cards:
            require(bool(c['exam']) == bool(c['exam_upgraded']) == bool(c['exam_condition']), c['id'])
            require(bool(c['chapter']) == bool(c['chapter_upgraded']), c['id'])
        require(all(c['owner'] == 'P02' and c['type'] == '策略' for c in chapters),
                'Chapters must be P02 strategies')
        require({c['id'] for c in tokens} == {'R3-T201', 'R3-T202'}, 'Token definitions')
        return {'exam_faces': 14, 'chapters': 6, 'tokens': 2}
    check('complete_exam_chapter_and_token_faces', faces)

    def starters():
        result = {}
        for owner, special in data['starters'].items():
            require(len(special) == 2 and len(set(special)) == 2, owner)
            require(all(cid in by_id and by_id[cid]['owner'] == owner
                        and by_id[cid]['permanent'] for cid in special), owner)
            eligible = [c for c in cards if c['owner'] in {owner, 'C'} and c['rarity'] != '基础']
            require(len(eligible) == 50, (owner, 'reward eligibility'))
            require(all(c['id'] not in {'R3-C-01', 'R3-C-02'} for c in eligible), owner)
            result[owner] = {'starter_copies': 4+4+1+1+len(special), 'eligible_ids': 50}
        require('固有' in by_id['R3-P02-01']['base'], 'Opening innate missing')
        require(by_id['R3-T201']['cost'] == 2 and not by_id['R3-T201']['permanent'], 'Publication lifecycle')
        return result
    check('starter_composition_and_reward_eligibility', starters)

    def md_rows():
        rows = {}
        for line in markdown.splitlines():
            if line.startswith('| R3-'):
                parts = [v.strip().replace('&#124;', '|').replace('<br>', '\n')
                         for v in line.strip('|').split('|')]
                if len(parts) == 8:
                    require(parts[0] not in rows, f'Duplicate card row {parts[0]}')
                    rows[parts[0]] = parts
        require(set(rows) == set(by_id), 'Markdown definition set differs')
        for cid, c in by_id.items():
            expected = [cid, c['name'], f"{c['cost']} → {c['upgraded_cost']}", c['type'],
                        c['rarity'], c['base'], c['upgraded'], '、'.join(c['tags'])]
            require(rows[cid] == expected, f'Markdown mismatch {cid}')
        for c in cards:
            if c['exam']:
                require(c['exam'] in markdown and c['exam_upgraded'] in markdown, c['id'])
            if c['chapter']:
                require(c['chapter'] in markdown and c['chapter_upgraded'] in markdown, c['id'])
        return {'matching_full_card_rows': len(rows)}
    check('markdown_all_card_fields_match_json', md_rows)

    def group_guides():
        require(len(guides) == 16, '16 guides required')
        require(Counter(g[2] for g in guides) == {'主流派': 11, '支持组': 4, '桥接组': 1},
                'Do not count support as independent routes')
        require(markdown.count('**设计思路**') == markdown.count('**典型联动**')
                == markdown.count('**特色与取舍**') == 16, 'Guide section coverage')
        for owner, title, role, idea, combo, tradeoff in guides:
            selected = [c for c in cards if c['owner'] == owner and c['group'] == title]
            require(len(selected) == 10, (owner, title))
            for value in (idea, combo, tradeoff):
                require(value in markdown, (owner, title, 'missing guide'))
            # The first card table must follow the group's introduction.
            heading = f'### {title}（10种，{role}）'
            guide_pos = markdown.index(idea, markdown.index(heading))
            require(guide_pos < markdown.index('| '+selected[0]['id']), (owner, title, 'guide position'))
        return {'guides': 16, 'main_routes': 11, 'support_or_bridge': 5}
    check('group_ideas_synergies_and_tradeoffs_precede_cards', group_guides)

    def workbook():
        expected_names = ['说明', '角色与被动', '卡牌总表', '林序', '许知微', '沈见章', '唐以宁',
                          '公共校园池', '分组导读', '章节常驻', '考试面', '论文衍生', '起手',
                          '负面状态', '功能覆盖']
        require(book.sheetnames == expected_names, book.sheetnames)
        mappings = {'卡牌总表': all_cards, '公共校园池': [c for c in cards if c['owner'] == 'C'],
                    '论文衍生': tokens}
        for owner in ('P01', 'P02', 'P03', 'P04'):
            mappings[data['owners'][owner]] = [c for c in cards if c['owner'] == owner]
        for sheet, content in mappings.items():
            actual = normalized(book[sheet].iter_rows(min_row=2, values_only=True))
            require(actual == [flat(c) for c in content], f'Workbook field mismatch: {sheet}')
        require(normalized(book['分组导读'].iter_rows(min_row=2, values_only=True)) == [list(g) for g in guides],
                'Workbook guides mismatch')
        expected_chapters = [[c[k] for k in ['id', 'name', 'chapter', 'chapter_upgraded']]
                             for c in cards if c['chapter']]
        require(normalized(book['章节常驻'].iter_rows(min_row=2, values_only=True)) == expected_chapters,
                'Workbook chapters mismatch')
        expected_exams = [[c[k] for k in ['id', 'name', 'exam_condition', 'type', 'cost',
                                          'upgraded_cost', 'exam', 'exam_upgraded']]
                          for c in cards if c['exam']]
        require(normalized(book['考试面'].iter_rows(min_row=2, values_only=True)) == expected_exams,
                'Workbook exam faces mismatch')
        require(normalized(book['负面状态'].iter_rows(min_row=2, values_only=True)) == data['debuffs'],
                'Workbook debuffs mismatch')
        expected_features = [[data['owners'][owner], feature, own, common]
                             for owner, feature, own, common in data['feature_coverage']]
        require(normalized(book['功能覆盖'].iter_rows(min_row=2, values_only=True)) == expected_features,
                'Workbook features mismatch')
        starter_rows = normalized(book['起手'].iter_rows(min_row=2, values_only=True))
        for owner in data['starters']:
            own = [row for row in starter_rows if row[0] == owner]
            require(sum(row[4] for row in own) == 12, owner)
            expected = {'R3-C-01': 4, 'R3-C-02': 4, 'R3-C-03': 1, 'R3-C-04': 1,
                        **{cid: 1 for cid in data['starters'][owner]}}
            require({row[2]: row[4] for row in own} == expected, owner)
        require(not any(cell.data_type == 'f' for ws in book for row in ws for cell in row),
                'Unexpected Excel formula')
        require(not book._external_links, 'Unexpected external links')
        for ws in book:
            require(ws.freeze_panes == 'B2' and ws.auto_filter.ref == ws.dimensions, ws.title)
        return {'sheets': 15, 'full_card_definitions': 174, 'roundtrip': True,
                'formulas': 0, 'external_links': 0}
    check('excel_readback_all_faces_guides_starters_and_safety', workbook)
    book.close()

    def debuff_contract():
        require(data['version'] == 'v3.1', 'Current design version must be v3.1')
        expected = {'侵蚀': 20, '诱发': 3, '脆化': 3, '乏力': 4, '虚弱': 3, '易伤': 3}
        require(len(data['debuffs']) == 6, 'Six distinct task debuffs required')
        require({row[0]: row[1] for row in data['debuffs']} == expected, 'Debuff names/caps')
        require(all(len(row) == 5 and all(str(v).strip() for v in row) for row in data['debuffs']),
                'Incomplete debuff metadata')
        own = [c for c in cards if c['owner'] == 'P03']
        require(Counter(c['group'] for c in own) == {
            '侵蚀增殖': 10, '诱发连打': 10, '混合抑制': 10, '通用衔接': 10}, 'P03 group identity')
        role = read('docs/角色与专属卡池重设计_v3.md')
        rules = read('docs/交大四年_DDL战争_游戏规则书.md')
        for text, label in ((role, 'role'), (rules, 'rules'), (markdown, 'card table')):
            for name, cap in expected.items():
                require(re.search(r'\| '+name+r' \| '+str(cap)+r'(?: \||；)', text),
                        (label, name, 'missing state/cap table'))
            require('复合反应' in text and '每目标每回合' in text and '4次' in text,
                    (label, 'missing passive or reaction limit'))
        for name in expected:
            require(any(name in c['base'] for c in own), (name, 'no actual P03 card source'))
        require('从0实际建立' in role and '单纯加厚原状态不符合' in role
                and '实际成功才消耗' in role, 'Passive candidate/capacity rules missing')
        require('刚建立的诱发不响应本次出牌' in role and '不减层' in role
                and '最多2/3次' in role, 'Reaction/erosion timing distinction missing')
        return {'debuff_caps': expected, 'p03_routes': 3, 'support_groups': 1,
                'scope': 'Written definitions and card sources only; behavior not executed'}
    check('six_debuff_definitions_passive_and_p03_route_contract', debuff_contract)

    def feature_coverage():
        required = {'抽牌', '防护', '群攻', '易伤', '虚弱', '破盾'}
        rows = data['feature_coverage']
        require(len(rows) == 24 and len({(r[0], r[1]) for r in rows}) == 24,
                'Four characters times six distinct functions required')
        patterns = {'抽牌': r'抽\d+', '防护': r'放松',
                    '群攻': r'(?:所有任务|其他当前任务).*?完成',
                    '易伤': r'易伤', '虚弱': r'虚弱', '破盾': r'移除(?:其|单任务)?至多\d+屏障'}
        for owner in data['starters']:
            require({r[1] for r in rows if r[0] == owner} == required, owner)
            for state in ('易伤', '虚弱'):
                require(any(c['owner'] == owner and state in c['base'] for c in cards),
                        (owner, state, 'exclusive cards missing common debuff function'))
        for owner, feature, own, common in rows:
            require(common in by_id and by_id[common]['owner'] == 'C', (owner, feature, common))
            require(own or feature == '破盾', (owner, feature, 'unexplained exclusive gap'))
            if own:
                require(own in by_id and by_id[own]['owner'] == owner, (owner, feature, own))
            for cid in filter(None, (own, common)):
                for face in ('base', 'upgraded'):
                    effect = by_id[cid][face]
                    if not by_id[cid]['permanent'] and face == 'upgraded':
                        require(effect == '不可升级；效果同基础面。', (cid, 'token upgrade rule'))
                        effect = by_id[cid]['base']
                    require(re.search(patterns[feature], effect),
                            (owner, feature, cid, face, 'source text does not provide function'))
        require({r[1]: r[3] for r in rows if r[0] == 'P01'} == {
            '抽牌': 'R3-C-03', '防护': 'R3-C-02', '群攻': 'R3-C-12',
            '易伤': 'R3-C-09', '虚弱': 'R3-C-10', '破盾': 'R3-C-11'}, 'Public function IDs')
        return {'character_function_pairs': 24, 'base_and_upgrade_source_texts_checked': True,
                'common_cards_added': 4, 'scope': 'Availability/definitions; acquisition not simulated'}
    check('all_character_common_functions_have_valid_card_sources', feature_coverage)

    def revision_card_changes():
        source = 'artifacts/redesign-20261030/campus-cards-v3.json'
        with zipfile.ZipFile(revision_dir / 'before-v3.1.zip') as archive:
            old_bytes = archive.read(source)
        old = json.loads(old_bytes.decode('utf-8'))
        before = {c['id']: c for c in old['cards'] + old['tokens']}
        added_ids = sorted(set(by_id) - set(before))
        removed_ids = sorted(set(before) - set(by_id))
        changes = []
        for cid in sorted(set(before) & set(by_id)):
            fields = {key: {'before': before[cid][key], 'after': by_id[cid][key]}
                      for key in FIELDS if before[cid][key] != by_id[cid][key]}
            if fields:
                changes.append({'id': cid, 'before_name': before[cid]['name'],
                                'after_name': by_id[cid]['name'], 'fields': fields})
        expected_modified = {f'R3-P03-{i:02}' for i in range(1, 41) if i != 34} | {
            'R3-P01-32', 'R3-P01-33', 'R3-P02-34', 'R3-P02-35', 'R3-P04-31', 'R3-P04-35'}
        require({c['id'] for c in changes} == expected_modified, 'Unexpected modified card set')
        require(added_ids == [f'R3-C-{i:02}' for i in range(9, 13)] and not removed_ids,
                {'added': added_ids, 'removed': removed_ids})
        require(old['tokens'] == tokens and old['owners'] == data['owners'], 'Token/owner changes out of scope')
        starter_changes = {owner: {'before': old['starters'][owner], 'after': new}
                           for owner, new in data['starters'].items() if old['starters'][owner] != new}
        require(starter_changes == {'P03': {'before': ['R3-P03-01', 'R3-P03-02'],
                                            'after': ['R3-P03-01', 'R3-P03-21']}}, 'Starter delta')
        output = {'from_version': old['version'], 'to_version': data['version'],
                  'before_archive': 'before-v3.1.zip', 'before_json_path': source,
                  'before_json_sha256': hashlib.sha256(old_bytes).hexdigest(),
                  'after_json_sha256': sha(HERE / 'campus-cards-v3.json'),
                  'modified_count': len(changes), 'added_count': len(added_ids),
                  'removed_count': len(removed_ids),
                  'unchanged_definition_count': len(before) - len(changes),
                  'p03_unchanged_support_card': 'R3-P03-34', 'modified': changes,
                  'added': [by_id[cid] for cid in added_ids], 'removed_ids': removed_ids,
                  'starter_changes': starter_changes}
        (revision_dir / 'card-changes.json').write_text(
            json.dumps(output, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        return {'modified': len(changes), 'added': len(added_ids), 'removed': 0,
                'unchanged_definitions': len(before)-len(changes), 'starter_changes': starter_changes}
    check('revision_card_delta_matches_authorized_scope', revision_card_changes)

    def reference_snapshot():
        reference = json.loads((revision_dir / 'reference-cards.json').read_text(encoding='utf-8'))
        ids = reference['selected_ids']
        require(len(ids) == 9 and len(set(ids)) == 9, 'Nine reference IDs required')
        require(set(ids) == {'DEADLY_POISON', 'POISONED_STAB', 'NOXIOUS_FUMES', 'BUBBLE_BUBBLE',
                             'ACCELERANT', 'CORROSIVE_WAVE', 'OUTBREAK', 'MALAISE', 'STRANGLE'},
                'Reference identity mismatch')
        require(reference['game_patch'] == 'unconfirmed'
                and reference['snapshot_date'] == '2026-10-29', 'Reference provenance boundary')
        for language in ('eng', 'zhs'):
            metadata = reference['sources'][language]
            path = ROOT / metadata['path']
            require(sha(path) == metadata['sha256'], (language, 'source hash changed'))
            source_rows = json.loads(path.read_text(encoding='utf-8'))
            require(len(source_rows) == metadata['rows'] == 577, (language, 'source count'))
            lookup = {r['id']: r for r in source_rows}
            require(reference['records'][language] == [lookup[cid] for cid in ids],
                    (language, 'reference record differs from archived original'))
        return {'reference_cards': 9, 'languages': 2, 'all_record_fields_and_source_hashes_match': True,
                'remote_sources_fetched_by_this_verifier': False, 'latest_patch_confirmed': False}
    check('nine_reference_cards_match_archived_bilingual_snapshots', reference_snapshot)

    documents = ([ROOT / 'README.md', ROOT / 'agent.md', ROOT / '实施计划与验收基线.md']
                 + sorted((ROOT / 'docs').rglob('*.md'))
                 + sorted((ROOT / '动作或者图片描述').rglob('*.md'))
                 + [HERE / 'README.md'])

    def links():
        broken, total = [], 0
        for path in documents:
            text = path.read_text(encoding='utf-8')
            require('\ufffd' not in text, f'Invalid UTF-8 replacement in {path.name}')
            text = re.sub(r'^```.*?^```[^\n]*', '', text, flags=re.M | re.S)
            text = re.sub(r'`[^`\n]+`', '', text)
            for match in re.finditer(r'\[[^\]\n]*\]\(([^\n]+?)\)', text):
                target = match.group(1).strip()
                if target.startswith('<') and '>' in target:
                    target = target[1:target.index('>')]
                else:
                    target = re.split(r'\s+["\']', target)[0]
                if re.match(r'^[A-Za-z][\w+.-]*:', target) or target.startswith(('#', '//')):
                    continue
                target = unquote(target.split('#', 1)[0])
                if not target:
                    continue
                total += 1
                if not (path.parent / target).exists():
                    broken.append({'source': path.relative_to(ROOT).as_posix(), 'target': target})
        require(not broken, broken)
        return {'markdown_files': len(documents), 'local_link_paths': total,
                'broken': [], 'scope': 'existing local paths; anchors and remote sites not fetched'}
    check('markdown_utf8_and_local_link_paths', links)

    def legacy():
        active = [p for p in documents if not p.name.startswith(('G0', 'G1', 'G2', 'G3', 'G4', 'G5', 'G6', 'G7'))]
        findings = []
        patterns = [r'当前设计以已改写的v2', r'新R2卡', r'R2-[CP]\d?\d?-\d+',
                    r'获得\d+(?:点)?(?:编译|数据|稿件)', r'(?:编译|稿件)(?:上限|初始)',
                    r'每人(?:均有)?(?:三|3)条(?:主)?流派']
        for path in active:
            for index, line in enumerate(path.read_text(encoding='utf-8').splitlines(), 1):
                if any(re.search(pattern, line) for pattern in patterns):
                    findings.append(f'{path.relative_to(ROOT)}:{index}: {line}')
        require(not findings, findings)
        for path in ['docs/角色与专属卡池重设计_v2.md', 'docs/校园卡牌设计表_v2.md',
                     '动作或者图片描述/角色卡池重设计需求_v2.md']:
            text = read(path)
            require('迁移入口' in text and 'v3' in text and len(text.splitlines()) <= 8, path)
        require(not (ROOT / 'docs/小怪设计草稿.md').exists(), 'Legacy draft still active')
        require('后续任务候选' in read('docs/交大四年_DDL战争_统一设计文档.md'), 'Draft ideas not consolidated')
        return {'patterns_checked': patterns, 'active_documents': len(active), 'findings': [],
                'boundary': 'Targeted text audit; not a semantic proof of all card rules'}
    check('targeted_legacy_rule_audit_and_consolidation', legacy)

    def retired_p03_rules():
        historical_names = ('G0', 'G1', 'G2', 'G3', 'G4', 'G5', 'G6', 'G7', '文档修改总结报告_')
        active = [p for p in documents if not p.name.startswith(historical_names)]
        patterns = [r'破绽|动摇|暴露|催化|引爆|毒爆', r'公共(?:8种|8张)',
                    r'168\s*(?:种)?永久', r'170个卡定义', r'40\+6', r'46种(?:奖励|资格)',
                    r'13张工作表', r'自然与额外(?:合计|触发)']
        findings = []
        for path in active:
            for index, line in enumerate(path.read_text(encoding='utf-8').splitlines(), 1):
                if any(re.search(pattern, line) for pattern in patterns):
                    findings.append(f'{path.relative_to(ROOT)}:{index}: {line}')
        require(not findings, findings)
        catalog = read('artifacts/redesign-20261030/catalog_v3.py')
        require(not re.search(patterns[0], catalog), 'Retired mechanics in current card source')
        return {'current_documents_scanned': len(active), 'patterns_checked': patterns,
                'findings': [], 'scope': 'Targeted current-rule audit; historical reports excluded'}
    check('retired_p03_mechanics_and_outdated_delivery_counts_removed', retired_p03_rules)

    def backup():
        with zipfile.ZipFile(HERE / 'before-v3.zip') as archive:
            require(archive.testzip() is None, 'Backup ZIP CRC failed')
            hashes = json.loads(archive.read('SHA256.json').decode('utf-8'))
            require(hashes == baseline['documents'], 'Backup manifest mismatch')
            for path, expected in hashes.items():
                actual = hashlib.sha256(archive.read(path)).hexdigest()
                require(actual == expected, f'Backup byte mismatch {path}')
        return {'files': len(baseline['documents']), 'crc_and_sha256_match': True}
    check('before_backup_crc_and_all_file_hashes', backup)

    def revision_backup():
        with zipfile.ZipFile(revision_dir / 'before-v3.1.zip') as archive:
            require(archive.testzip() is None, 'Revision ZIP CRC failed')
            hashes = json.loads(archive.read('SHA256.json').decode('utf-8'))
            require(hashes == revision_before and len(hashes) == 56, 'Revision backup manifest mismatch')
            require(set(archive.namelist()) == set(hashes) | {'SHA256.json'}, 'Unexpected ZIP entries')
            for path, expected in hashes.items():
                require(hashlib.sha256(archive.read(path)).hexdigest() == expected,
                        f'Revision backup byte mismatch: {path}')
        return {'files': len(revision_before), 'crc_and_all_sha256_match': True}
    check('revision_backup_crc_and_all_56_file_hashes', revision_backup)

    def protected_files():
        errors = [p for p, expected in protected.items()
                  if not (ROOT / p).is_file() or sha(ROOT / p) != expected]
        for folder in ('src', 'data', 'tests', 'tools', 'assets'):
            current = {p.relative_to(ROOT).as_posix() for p in (ROOT / folder).rglob('*') if p.is_file()}
            expected = {p for p in protected if p.startswith(folder + '/')}
            errors.extend(sorted(current.symmetric_difference(expected)))
        require(not errors, errors)
        historical = [p for p in baseline['documents']
                      if re.match(r'docs/G[0-7]', p) or p == 'AGENTS.md']
        for path in historical:
            require(sha(ROOT / path) == baseline['documents'][path], f'Historical evidence changed: {path}')
        return {'protected_files': len(protected), 'unchanged_historical_records_and_instructions': len(historical),
                'game_engine_executed': False, 'player_save_access': False}
    check('program_data_assets_tests_tools_research_and_evidence_unchanged', protected_files)

    def boundaries():
        for path in ['README.md', '实施计划与验收基线.md', 'docs/四角色与卡池重设计汇总_20261030.md',
                     'docs/角色与专属卡池重设计_v3.md']:
            text = read(path)
            require('v1' in text and 'G0' in text and 'G8' in text and '未' in text, path)
        report = read('docs/四角色与卡池重设计汇总_20261030.md')
        require('两条真正改变收益方式' in report and '中间牌不直接产生完成量' in report,
                'Publication/chapter design change is not disclosed')
        revision = read('docs/沈见章毒师与全角色通用功能修订_20261030.md')
        require('v3.1' in revision and '六负面' in revision and '39张' in revision
                and '45个既有定义' in revision and '新增4个公共定义' in revision,
                'Revision scope/delta not disclosed')
        require('第三方API快照' in revision and '游戏补丁未确认' in revision
                and '不等于原作的第二种Poison' in revision and '本项目扩展' in revision,
                'Original/reference distinction not disclosed')
        require('未运行Godot' in revision and '全部未执行' in revision and 'G8未开始' in revision,
                'Revision unverified boundary missing')
        return {'design_only': True, 'future_behavior_cases_executed': 0,
                'g0_or_g8_claimed_passed_for_v3': False}
    check('design_implementation_and_acceptance_boundaries', boundaries)

    changed = []
    for path, expected in baseline['documents'].items():
        file = ROOT / path
        status = 'deleted' if not file.exists() else ('modified' if sha(file) != expected else 'unchanged')
        if status != 'unchanged':
            changed.append({'path': path, 'status': status})
    revision_changed = []
    for path, expected in revision_before.items():
        file = ROOT / path
        status = 'deleted' if not file.exists() else ('modified' if sha(file) != expected else 'unchanged')
        if status != 'unchanged':
            revision_changed.append({'path': path, 'status': status})
    exit_code = int(any(not r['passed'] for r in RESULTS))
    snapshots = {p.relative_to(ROOT).as_posix(): sha(p) for p in documents if p.exists()}
    for path in HERE.iterdir():
        if path.is_file() and path.name not in {'validation.json', 'validation.md'}:
            snapshots[path.relative_to(ROOT).as_posix()] = sha(path)
    for path in revision_dir.rglob('*'):
        if path.is_file():
            snapshots[path.relative_to(ROOT).as_posix()] = sha(path)
    report = {'timestamp': datetime.now().astimezone().isoformat(), 'command': COMMAND,
              'environment': {'python': sys.version, 'os': platform.platform(),
                              'openpyxl': openpyxl.__version__,
                              'engine': 'not invoked; design document verification only'},
              'exit_code': exit_code, 'checks': RESULTS, 'document_changes': changed,
              'revision_changes': revision_changed,
              'reviewed_file_sha256': snapshots,
              'not_verified': ['Game behavior or balance', 'Godot G0-G8 acceptance',
                               'Future R3-T01-23 behavior cases', 'Excel desktop visual review',
                               'Remote sources/latest game patch', 'Exhaustive loop/combo analysis']}
    (HERE / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    lines = ['# v3.1设计交付静态核查记录', '', f"时间：{report['timestamp']}", '',
             f'命令：`{COMMAND}`', '', f'退出码：{exit_code}', '',
             f"环境：Python {platform.python_version()}；{platform.platform()}；openpyxl {openpyxl.__version__}；未启动Godot。", '',
             '| 检查 | 实际结果 |', '|---|---|']
    lines += [f"| {r['check']} | {'通过' if r['passed'] else '失败'} |" for r in RESULTS]
    lines += ['', '## 实际输出', '', '```text']
    for r in RESULTS:
        lines.append(('PASS ' if r['passed'] else 'FAIL ')+r['check'])
        if not r['passed']:
            lines.append(str(r['detail']))
    lines += [f'EXIT_CODE={exit_code}', '```', '', '## 修改及删并清单', '',
              '| 路径 | 与改前备份相比 |', '|---|---|']
    lines += [f"| {r['path']} | {r['status']} |" for r in changed]
    lines += ['', '## v3.1本轮差异', '', '| 路径 | 与before-v3.1相比 |', '|---|---|']
    lines += [f"| {r['path']} | {r['status']} |" for r in revision_changed]
    lines += ['', '本轮新增[毒师修订报告](../../docs/沈见章毒师与全角色通用功能修订_20261030.md)、[逐牌差异](debuff-revision/card-changes.json)及[双语参考摘录](debuff-revision/reference-cards.json)。三者随当前报告分别保存，不回写改前备份。']
    lines += ['', '新增交付为角色v3、卡表v3、[新汇总](../../docs/四角色与卡池重设计汇总_20261030.md)、v3文字需求及本目录内设计表格/数据/核查文件。完整字段与来源SHA256见[机器报告](validation.json)。', '',
              '这是文档、设计数据与Excel的一致性证据。未测试游戏行为、卡组平衡、无限组合、未来R3-T01～23用例、G0/G8、桌面Excel显示或远程来源。历史研究数据和现有程序仍对应原版本；没有读取/修改玩家存档。', '']
    (HERE / 'validation.md').write_text('\n'.join(lines), encoding='utf-8')
    print(f'EXIT_CODE={exit_code}')
    return exit_code


if __name__ == '__main__':
    raise SystemExit(main())
