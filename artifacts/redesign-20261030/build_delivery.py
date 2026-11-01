"""Build review documents and workbook; not a game content pipeline."""
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DATA = json.loads((HERE / 'campus-cards-v3.json').read_text(encoding='utf-8'))
CARDS, TOKENS = DATA['cards'], DATA['tokens']
CHARACTERS = [
    ('P01', '林序', '步骤复用', '首个相邻策略→工作结算后，可取回前一永久策略；每回合一次，不减费。', '交替连携、策略复用、集中执行'),
    ('P02', '许知微', '论文记录', '开题后首次牌名加进度，重复牌名加引用；同名每回合最多打一次，后三级引用额外付专注。', '发表构筑、章节引擎；另有路线衔接与通用支持'),
    ('P03', '沈见章', '复合反应', '每回合首次由卡牌正文给已有另一种负面的任务新增一种负面后，该新状态额外+1层；每回合一次，不递归。', '侵蚀增殖、诱发连打、混合抑制'),
    ('P04', '唐以宁', '分工默契', '每回合首张合作结算后获得1协同、2放松，协同上限6；考试面仍触发。', '分工推进、接力轮换、备考内化'),
]
GUIDES = [('P01', '交替连携', '主流派', '用策略→工作交替出牌，让低费工作获得额外完成量、放松或抽牌。构筑优先考虑费用曲线和策略/工作比例。', '快速定位找牌→小步实践触发连携并取回快速定位；后续用检查输入→分步运行继续连携，一气呵成按此前连携次数收尾。', '收益取决于相邻顺序，能力或场景会打断；持续多次小输出，受专注和手牌数量约束。'), ('P01', '策略复用', '主流派', '反复使用同一张高价值策略，形成防护和过牌循环；核心是让一张牌多次发挥功能。', '断点休息→复用验证，借被动取回断点休息；以后再次使用同实例，让结果核验和方法沉淀兑现复用收益。', '选牌重视值得取回的策略，输出密度可低于连携路线；回收不减费，同一实例成长只到第二次固定增强。'), ('P01', '集中执行', '主流派', '把数张准备牌的收益集中到一张重费工作，构成准备→爆发的回合。', '列执行步骤叠首段加值，预留时间减费，再用集中运行；多目标时用整体验收，归档重跑为下一轮取回。', '区别于多次连携，它追求少量高质量输出；准备总加值上限12，群体加值只给主目标，坏起手时可能准备不足。'), ('P01', '通用衔接', '支持组', '补齐过牌、防护、保留和基础负面；群体输出由接口对齐、整体验收及公共牌提供。', '拆分需求用易伤衔接工作，保存现场用虚弱承接策略复用；桌前整理找牌，交接备注保留重费工作。', '通用支持不单列路线；易伤虚弱与其他人共享，林序特色仍是顺序与复用。破盾可由公共扫清障碍补位。'), ('P02', '发表构筑', '主流派', '用不同牌名扩大覆盖率，再用引用与修正把论文推到合适完成度，最后一次发表兑现完成量。', '开题→检索目录、阅读摘记等不同名字→交叉引证和定稿检查补足→解题发表；DDL紧时用先交可用版提前交付。', '选牌偏好多样牌名、过牌和防护；收益集中在发表，满100有额外奖励，但等待会占用DDL，发表后章节全部退场。'), ('P02', '章节引擎', '主流派', '部署最多三章，让它们从下回合起按槽序自动运行，依靠持续抽牌、防护和完成量推进任务。', '引言章提供过牌，方法章维持防护，结论章放在后槽读取此前运行数；也可用结果章作稳定输出，集中研讨提供有限额外激活。', '构筑优先章节检索、槽序和存活，不以攒满论文为主要收益；启动较慢，而且发表会拆掉整个引擎。'), ('P02', '路线衔接', '桥接组', '帮助发表与章节两条路线切换，并把同名牌、增幅和修正纳入决策。它没有单独的胜利闭环。', '章节回看把牌留到下轮再用获得引用；综述章的增幅配合页间留白，阶段总结撤章换修正后转发表。', '同名循环只在跨回合成立；换题不刷新回合限次，引用转述会失去原引用的6点贡献，换抽牌和较低引用付款档位。'), ('P02', '通用支持', '支持组', '提供生存、抽牌、专注和再开题入口，并用易伤、虚弱补齐直接输出与压力控制。', '预留余量施虚弱并承接引用付款，独立推演施易伤后完成任务；发表后用收拢材料取回开题，公共破盾和群攻补齐应急工具。', '透支直接失精力且耗尽；章节自动输出不吃易伤，解题发表属于工作牌可获益。基础负面不创造新的论文资源路线。'), ('P03', '侵蚀增殖', '主流派', '把侵蚀当作持续毒：先建立层数，再用条件施毒、增殖与结算频率放大输出。牌组用较少出牌次数换取后续任务阶段的持续收益。', '圈出疑点或破题追问施毒→疑点滋生加厚→提前验算提前结算；追踪记录稳定施毒，加速剖析增加自然结算次数。留痕比对补虚弱，触发复合反应并争取轮数。', '主要选施毒、增殖和防护；不消耗毒层换爆发。短DDL或阶段替换会压缩收益，额外自然结算每次也减1层，扩散需要第二目标。'), ('P03', '诱发连打', '主流派', '把诱发当作反应毒：先挂在目标上，再用低费牌与过牌连续触发，当前回合就获得收益。检索连锁将抽牌转为施毒，连接持续路线。', '设下追问挂诱发→短句反问、翻页核证、逐句质询接连出牌；交叉污染先给诱发再给易伤可强化新状态，二次追问提供有限额外触发。检索连锁与循证爆发奖励实际抽牌和施毒次数。', '主要选低费、抽牌和当回合可执行的顺序，诱发行动结束清除，每目标每回合最多4次；重新挂诱发不刷新次数。不靠攒厚侵蚀，也不要求稀有能力才能运转。'), ('P03', '混合抑制', '主流派', '让不同负面承担不同工作：虚弱和乏力降低压力，脆化减少新屏障，易伤强化工作输出。以负面种类数支持直接进攻，并靠混合状态维持续航。', '削弱论势施虚弱和乏力，拆解防线腐蚀防御，再用公开前提施易伤；论点交锋读取种类数，全面质询处理多目标。复合论证可快速建立不同状态。', '主要选多种负面、工作输出和防护，少选单纯增殖毒层；论点交锋最多计4种且不消费状态。基础易伤虚弱并非专属，特色来自混合效率和被动强化。'), ('P03', '通用衔接', '支持组', '提供抽牌、防护、纠错、回收和保留，帮助三种收益方式在同一牌组中混搭；不新增第四条路线。', '整理思绪找启动牌，留待答辩保留配方；澄清误会解除焦虑并补虚弱，复盘取证为连打筛牌，配方笔记承接复合反应。', '工具要与输出比例平衡。回收仍需支付费用且不刷新限次；没有负面新目标可由直陈要点先处理。'), ('P04', '分工推进', '主流派', '建立多伙伴场面，工作牌读取在场人数，并用有限指挥提升当回合推进。', '室友搭把手召入伙伴→分工清单备协同→各做一题；协调推进或全组研讨让伙伴有限额外行动。', '选牌重视伙伴数量和工作收益；额外行动也消耗伙伴精力，不能只顾爆发忽略驻场时间。'), ('P04', '接力轮换', '主流派', '恢复伙伴精力、保住重要支援，并在短驻伙伴退出前安排交接，获得稳定防护和过牌。', '临时补位提供短期支援，轮换排班和轮值支援延长驻场；需要空槽时让同学休息，交接笔记接续手牌。', '重视续航和合作触发，推进节奏较稳；主动撤回不退卡、不触发额外离场收益，额外行动会加快精力消耗。'), ('P04', '备考内化', '主流派', '少依赖伙伴数量，使用协同门槛、无单位收益及合作考试面，建立普通战与考试通用的个人输出。', '分工内化或独立演练启动协同→先做会的、分工已内化保持门槛收益→独立交卷清空协同终结。', '主要收益来自主角手牌而非自动伙伴；保留协同可维持门槛，花掉则换当前收益，考试仍正常识别合作类型。'), ('P04', '通用衔接', '支持组', '补齐协同启动、筛选、防护、群攻及基础易伤虚弱，让三条路线都能应对常见任务。', '约定进度施虚弱并启动协同；各抒己见先给主目标易伤再群体输出，保留意见留终结牌。公共扫清障碍提供直接破盾。', '易伤只强化工作牌，不能放大伙伴自动协作；伙伴身份仍由分工、驻场和轮换体现。支持组不是第四条路线。')]

def table(headers, rows):
    def esc(v):
        return str(v).replace('|', '&#124;').replace('\n', '<br>')
    return '\n'.join(['| '+' | '.join(headers)+' |', '|'+ '|'.join(['---']*len(headers))+'|'] + ['| '+' | '.join(esc(v) for v in row)+' |' for row in rows])

def card_rows(cards):
    return [[c['id'], c['name'], f"{c['cost']} → {c['upgraded_cost']}", c['type'], c['rarity'], c['base'], c['upgraded'], '、'.join(c['tags'])] for c in cards]

def markdown():
    text = '''# 校园卡牌设计表 v3.1

日期：2026-10-30｜状态：设计提案；未接入程序、未做游戏平衡实测。

本表包含172种永久牌（公共12种、四人各40种），以及2种论文衍生定义。升级面、考试面和章节复印的临时实例不另计卡种；遭遇注入的X01负担另属任务原型。四人共11条主流派、16个功能分组；许知微只有发表与章节两条主路线。

机制依据为[角色设计](角色与专属卡池重设计_v3.md)与[规则书](交大四年_DDL战争_游戏规则书.md)，逐牌数据维护在[设计JSON](../artifacts/redesign-20261030/campus-cards-v3.json)，[Excel](../artifacts/redesign-20261030/校园角色与专属卡池_v3.xlsx)按同一数据生成。组前说明解释选牌、联动和取舍，不增加卡文以外的效果。[毒师与通用功能修订](沈见章毒师与全角色通用功能修订_20261030.md)记录替换范围与核查。

## 1. 阅读约定

“完成”减少任务未完成量、先经过屏障；“放松”吸收压力；“抽n”按容量逐张抽；“回收”只取回合法永久牌且不减费。“连携”指前一张实际打出的是策略、当前是工作。“引擎”指投入后持续带来收益的组合。“侵蚀”是任务阶段持续毒，“诱发”是出牌触发的当回合反应毒；“脆化/乏力”分别削防与减压，“易伤/虚弱”是全角色通用功能。状态完成量不吃工作加值或易伤，不能混算。

费用列为基础→升级；升级正文为完整替代文本。未写群体的完成量作用一个合法任务；当前工作首段加值群体只给主目标。每人普通20、进阶15、稀有5。牌组可在同人物内混搭，不需要凑齐某组10张。

## 2. 公共校园池（12种）

公共牌提供基本完成量、过牌、防护、恢复，以及易伤、虚弱、破盾、群攻，专属机制仍由人物卡和被动承担。写作业、伸个懒腰是基础牌，不进奖励；其余10种与该人物40种构成50种奖励资格。

'''
    headers=['ID','牌名','专注费用','类型','稀有度','基础完整效果','升级完整效果','功能标签']
    text += table(headers, card_rows([c for c in CARDS if c['owner']=='C']))+'\n'
    for index,(owner,name,passive,desc,routes) in enumerate(CHARACTERS,3):
        text += f'\n## {index}. {name}（40种）\n\n被动 **{passive}**：{desc} 详细边界以角色设计为准。\n'
        for _, title, role, idea, combo, tradeoff in [g for g in GUIDES if g[0]==owner]:
            text += f'\n### {title}（10种，{role}）\n\n**设计思路**：{idea}\n\n**典型联动**：{combo}\n\n**特色与取舍**：{tradeoff}\n\n'
            text += table(headers,card_rows([c for c in CARDS if c['owner']==owner and c['group']==title]))+'\n'
    text += '\n## 7. 六章常驻效果\n\n最多3个有序槽，部署后下回合才生效。下表是自动运行内容；手动从论文区使用只结算卡牌基础效果及记账。集中研讨另开额外窗口，遵守每章每回合最多一次额外激活。\n\n'
    text += table(['ID','章节','基础常驻','升级常驻'],[[c['id'],c['name'],c['chapter'],c['chapter_upgraded']] for c in CARDS if c['chapter']])+'\n'
    text += '\n## 8. 完整考试面（14种）\n\n先替换正文再校验目标与费用，保留主类型。固定考场只替换2张场景；独立作答替换7张合作与5张伙伴依赖策略。章节不被考试禁用。下列费用与普通面相同，正文不叠加执行。\n\n'
    text += table(['ID','牌名','条件','类型','费用基础→升级','基础考试全文','升级考试全文'],[[c['id'],c['name'],c['exam_condition'],c['type'],f"{c['cost']} → {c['upgraded_cost']}",c['exam'],c['exam_upgraded']] for c in CARDS if c['exam']])+'\n'
    text += '\n## 9. 论文衍生（不计永久数量）\n\n开题生成/取回同一个发表实例；旁听笔记统一同名。章节复印沿用源章定义ID，实例附临时、不可部署、使用耗尽标记，不创造额外牌名或永久资格。\n\n'
    text += table(headers, card_rows(TOKENS))+'\n'
    text += '\n## 10. 维护与验证边界\n\n修改逐牌数据后须同步生成Markdown与Excel。现有v1数据、旧研究Excel、旧校园v2 Excel及G0历史登记保持原字节。程序行为、路线强弱、任务数值和所有R3-T01～23未来用例均未实测；文档一致性核查不替代新版G0或G8验收。\n'
    text += '\n## 11. 六种任务负面\n\n下表统一所有人物使用的状态语义，易伤与虚弱没有角色独占。详细时序与复合反应见角色设计§5。\n\n'
    text += table(['状态','上限','作用','到期/限制','定位'], DATA['debuffs'])+'\n'
    text += '\n## 12. 全角色通用功能覆盖\n\n“专属来源”可以是本人物卡或专有衍生，“公共补位”供四人获取。空白表示由公共牌补位，不代表缺失牌文。工作牌本就可消耗屏障，此处破盾专指明确移除屏障。\n\n'
    def label(cid):
        return cid+' '+next(c['name'] for c in CARDS+TOKENS if c['id']==cid) if cid else '—'
    text += table(['人物','功能','专属来源','公共补位'], [[DATA['owners'][o],f,label(a),label(b)] for o,f,a,b in DATA['feature_coverage']])+'\n'
    (ROOT/'docs/校园卡牌设计表_v3.md').write_text(text,encoding='utf-8')

FIELDS=['id','owner','character','name','name_key','group','cost','upgraded_cost','type','rarity','base','upgraded','tags','permanent','exam_condition','exam','exam_upgraded','chapter','chapter_upgraded']
HEADERS=['ID','归属','角色','牌名','语义牌名键','功能分组','基础费','升级费','类型','稀有度','基础完整效果','升级完整效果','标签','永久牌','考试条件','基础考试面','升级考试面','基础章节常驻','升级章节常驻']

def flat(c):
    return ['、'.join(c[k]) if isinstance(c[k],list) else c[k] for k in FIELDS]

def workbook():
    wb=Workbook(); wb.remove(wb.active)
    sheet_index=0
    def sheet(name,headers,rows,widths=None):
        nonlocal sheet_index
        sheet_index+=1
        ws=wb.create_sheet(name); ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
        ws.sheet_view.zoomScale=85
        ws.row_dimensions[1].height=30
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='224B56'); cell.font=Font(color='FFFFFF',bold=True)
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=105 if len(headers)>10 else 72
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.row%2==0: cell.fill=PatternFill('solid',fgColor='EFF5F4')
        from openpyxl.utils import get_column_letter
        for i,h in enumerate(headers,1):
            width=(widths[i-1] if widths else (64 if any(s in h for s in ['效果','考试面','常驻']) else (24 if h in ['标签','功能分组'] else 18)))
            ws.column_dimensions[get_column_letter(i)].width=width
        if ws.max_row>1:
            tab=Table(displayName=f'Review{sheet_index}',ref=ws.dimensions)
            tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(tab)
        return ws
    sheet('说明',['项目','说明'],[
        ['版本','v3.1 / 2026-10-30 / 文档提案，未实现与平衡实测'],
        ['数量','172永久＝公共12＋四人各40；2论文衍生，不以升级/考试面充数'],
        ['路线','全体11条主流派、16功能分组；P02只有发表、章节两主路线'],
        ['权威','角色与专属卡池重设计_v3.md维护机制；规则书维护共同顺序；campus-cards-v3.json维护逐牌数据'],
        ['研究','杀戮尖塔2研究表另在20261029目录，第三方快照、游戏补丁未确认；本表为校园原创方案'],
        ['操作','各表首行筛选、冻结首行及首列、全文换行；总表含全部基础/升级/考试/章节字段'],
        ['边界','现有v1程序与数据未修改；Excel未在桌面逐页人工验收；回读验证见validation.json'],
    ],[22,120])
    sheet('角色与被动',['角色ID','人物','被动','规则概要','主流派'],CHARACTERS,[18,18,20,85,64])
    sheet('卡牌总表',HEADERS,[flat(c) for c in CARDS+TOKENS])
    for owner,name,*_ in CHARACTERS:
        sheet(name,HEADERS,[flat(c) for c in CARDS if c['owner']==owner])
    sheet('公共校园池',HEADERS,[flat(c) for c in CARDS if c['owner']=='C'])
    sheet('分组导读',['归属','分组','性质','设计思路','典型联动','特色与取舍'],GUIDES,[16,20,16,88,88,88])
    sheet('章节常驻',['ID','章节','基础常驻','升级常驻'],[[c['id'],c['name'],c['chapter'],c['chapter_upgraded']] for c in CARDS if c['chapter']],[20,20,94,94])
    sheet('考试面',['ID','牌名','触发条件','主类型','基础费','升级费','基础考试全文','升级考试全文'],[[c[k] for k in ['id','name','exam_condition','type','cost','upgraded_cost','exam','exam_upgraded']] for c in CARDS if c['exam']],[20,20,18,14,14,14,88,88])
    sheet('论文衍生',HEADERS,[flat(c) for c in TOKENS])
    starters=[]
    for owner,name,*_ in CHARACTERS:
        for cid,num in [('R3-C-01',4),('R3-C-02',4),('R3-C-03',1),('R3-C-04',1)]+[(cid,1) for cid in DATA['starters'][owner]]:
            card=next(c for c in CARDS if c['id']==cid)
            starters.append([owner,name,cid,card['name'],num,'固有，占首次抽牌一格' if cid=='R3-P02-01' else '永久起手'])
    sheet('起手',['归属','人物','卡ID','牌名','数量','说明'],starters,[16,16,22,22,14,60])
    sheet('负面状态',['状态','上限','作用','到期与限制','定位'],DATA['debuffs'],[18,12,85,85,36])
    sheet('功能覆盖',['人物','功能','专属来源ID','公共补位ID'],[[DATA['owners'][o],f,a,b] for o,f,a,b in DATA['feature_coverage']],[20,18,28,28])
    path=HERE/'校园角色与专属卡池_v3.xlsx'; wb.save(path)
    check=load_workbook(path,data_only=False)
    actual=list(check['卡牌总表'].values)
    expected=[HEADERS]+[flat(c) for c in CARDS+TOKENS]
    # Excel empty cells return None; JSON deliberately uses empty strings.
    normalized=[[v if v is not None else '' for v in row] for row in actual]
    assert normalized==expected, 'Workbook field roundtrip mismatch'
    assert not any(cell.data_type=='f' for ws in check for row in ws for cell in row), 'Unexpected formula'
    check.close()
    return {'sheets':len(wb.sheetnames),'permanent':len(CARDS),'tokens':len(TOKENS),'field_roundtrip':True}

if __name__=='__main__':
    markdown()
    print(json.dumps(workbook(),ensure_ascii=False))
